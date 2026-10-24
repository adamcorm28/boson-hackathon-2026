"""Build an .xlsx workbook from a JSON spec. Formulas stay live in Excel.

Usage: python make_model.py <spec.json> <out.xlsx>
Spec:  {"sheets": [{"name": str, "rows": [[cell, ...], ...], "formulas": {"B5": "=SUM(B2:B4)"}?}, ...]}
The first row of each sheet is treated as a header and set in bold.
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def validate(spec: dict) -> None:
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ValueError('spec needs a non-empty "sheets" list')
    for i, sheet in enumerate(sheets):
        if not isinstance(sheet, dict) or not isinstance(sheet.get("name"), str):
            raise ValueError(f'sheets[{i}] needs a string "name"')
        if not isinstance(sheet.get("rows", []), list) or not isinstance(sheet.get("formulas", {}), dict):
            raise ValueError(f'sheets[{i}]: "rows" must be a list and "formulas" an object')


def build(spec: dict, out_path: str) -> None:
    validate(spec)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in spec["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        for row in sheet.get("rows", []):
            ws.append(row)
        for ref, formula in sheet.get("formulas", {}).items():
            ws[ref] = formula
        for cell in ws[1] if ws.max_row else []:
            cell.font = Font(bold=True)
        for col in range(1, ws.max_column + 1):
            width = max((len(str(c.value)) for c in ws[get_column_letter(col)] if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 50)
    wb.save(out_path)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(__doc__, file=sys.stderr)
        return 2
    try:
        with open(argv[1]) as f:
            build(json.load(f), argv[2])
    except (ValueError, json.JSONDecodeError) as e:
        print(f"spec error: {e}", file=sys.stderr)
        return 2
    print(argv[2])
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
