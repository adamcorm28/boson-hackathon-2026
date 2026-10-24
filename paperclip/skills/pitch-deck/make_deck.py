"""Build a .pptx deck from a JSON spec.

Usage: python make_deck.py <spec.json> <out.pptx>
Spec:  {"title": str, "subtitle": str?, "slides": [{"title": str, "bullets": [str, ...]}, ...]}
"""

import json
import sys

from pptx import Presentation

TITLE_LAYOUT, CONTENT_LAYOUT = 0, 1


def validate(spec: dict) -> None:
    if not isinstance(spec.get("title"), str):
        raise ValueError('spec needs a string "title"')
    if not isinstance(spec.get("slides"), list):
        raise ValueError('spec needs a "slides" list')
    for i, slide in enumerate(spec["slides"]):
        if not isinstance(slide, dict) or not isinstance(slide.get("title"), str):
            raise ValueError(f'slides[{i}] needs a string "title"')
        if not all(isinstance(b, str) for b in slide.get("bullets", [])):
            raise ValueError(f'slides[{i}].bullets must be a list of strings')


def build(spec: dict, out_path: str) -> None:
    validate(spec)
    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[TITLE_LAYOUT])
    cover.shapes.title.text = spec["title"]
    cover.placeholders[1].text = spec.get("subtitle", "")
    for slide_spec in spec["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[CONTENT_LAYOUT])
        slide.shapes.title.text = slide_spec["title"]
        body = slide.placeholders[1].text_frame
        bullets = slide_spec.get("bullets", [])
        body.text = bullets[0] if bullets else ""
        for bullet in bullets[1:]:
            body.add_paragraph().text = bullet
    prs.save(out_path)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(__doc__, file=sys.stderr)
        return 2
    try:
        with open(argv[1]) as f:
            build(json.load(f), argv[2])
    except (ValueError, json.JSONDecodeError) as e:
        print(f"spec error: {e}", file=sys.stderr)
        return 2
    print(argv[2])
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
