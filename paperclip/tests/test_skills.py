import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import pptx
import pytest

PAPERCLIP_DIR = Path(__file__).parent.parent
SKILLS = PAPERCLIP_DIR / "skills"
AGENTS = PAPERCLIP_DIR / "agents"
ROLE_IDS = ("ceo", "cfo", "eng_lead", "engineer", "sales_lead")
# Filled in by the seed script (B7) when it uploads each AGENTS.md to Paperclip.
PLACEHOLDERS = {"ARTIFACTS_DIR", "GATEWAY_PUBLIC_URL", "SKILLS_DIR", "PYTHON", "DEMO_REPO", "COMPANY_DATA_DIR"}
RESULT_CONVENTION = (
    "RESULT: <2–4 sentence summary>",
    "ARTIFACT: <kind> <url> <label>",
)


def run(script: Path, *args) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(script), *map(str, args)], capture_output=True, text=True)


def write_json(path: Path, data) -> Path:
    path.write_text(json.dumps(data))
    return path


# --- make_deck.py -----------------------------------------------------------------------

DECK = SKILLS / "pitch-deck" / "make_deck.py"


def test_make_deck_slide_count(tmp_path):
    spec = write_json(tmp_path / "deck.json", {
        "title": "Scranton Prospects", "subtitle": "Dunder Mifflin",
        "slides": [{"title": f"Slide {i}", "bullets": ["a", "b"]} for i in range(3)],
    })
    out = tmp_path / "deck.pptx"
    result = run(DECK, spec, out)
    assert result.returncode == 0, result.stderr
    assert len(pptx.Presentation(out).slides) == 4


def test_make_deck_contains_text(tmp_path):
    spec = write_json(tmp_path / "deck.json", {
        "title": "Scranton Prospects", "subtitle": "Q4",
        "slides": [{"title": "Keystone Community College", "bullets": ["$42k/yr", "Three campuses"]}],
    })
    out = tmp_path / "deck.pptx"
    run(DECK, spec, out)
    text = " ".join(
        shape.text_frame.text
        for slide in pptx.Presentation(out).slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "Scranton Prospects" in text and "Keystone Community College" in text and "Three campuses" in text


def test_make_deck_bad_spec_fails_with_message(tmp_path):
    spec = write_json(tmp_path / "deck.json", {"slides": "nope"})
    result = run(DECK, spec, tmp_path / "deck.pptx")
    assert result.returncode != 0
    assert "title" in result.stderr


# --- make_model.py ----------------------------------------------------------------------

MODEL = SKILLS / "excel-model" / "make_model.py"


def test_make_model_formula_preserved(tmp_path):
    spec = write_json(tmp_path / "model.json", {"sheets": [{
        "name": "Lackawanna",
        "rows": [["Item", "Value"], ["Year 1", 120000], ["Year 2", 120000], ["Year 3", 120000]],
        "formulas": {"B5": "=SUM(B2:B4)"},
    }]})
    out = tmp_path / "model.xlsx"
    result = run(MODEL, spec, out)
    assert result.returncode == 0, result.stderr
    ws = openpyxl.load_workbook(out)["Lackawanna"]
    assert ws["B5"].value == "=SUM(B2:B4)"
    assert ws["B2"].value == 120000


def test_make_model_multiple_sheets(tmp_path):
    spec = write_json(tmp_path / "model.json", {"sheets": [
        {"name": "Summary", "rows": [["a"]]}, {"name": "Detail", "rows": [["b"]]}]})
    out = tmp_path / "model.xlsx"
    run(MODEL, spec, out)
    assert openpyxl.load_workbook(out).sheetnames == ["Summary", "Detail"]


def test_make_model_bad_spec_fails_with_message(tmp_path):
    spec = write_json(tmp_path / "model.json", {"sheets": [{"rows": []}]})
    result = run(MODEL, spec, tmp_path / "model.xlsx")
    assert result.returncode != 0
    assert "name" in result.stderr


# --- AGENTS.md ----------------------------------------------------------------------------


@pytest.fixture(scope="module")
def agents_md():
    return {role: (AGENTS / role / "AGENTS.md").read_text() for role in ROLE_IDS}


def test_agents_md_have_result_convention(agents_md):
    for role, text in agents_md.items():
        for line in RESULT_CONVENTION:
            assert line in text, (role, line)


def test_agents_md_artifact_rule_and_time_limit(agents_md):
    for role, text in agents_md.items():
        assert "{ARTIFACTS_DIR}" in text, role
        assert "{GATEWAY_PUBLIC_URL}/artifacts/" in text, role
        assert "finish in under 3 minutes" in text, role


def test_agents_md_only_known_placeholders(agents_md):
    for role, text in agents_md.items():
        assert set(re.findall(r"\{([A-Z_]+)\}", text)) <= PLACEHOLDERS, role


def test_agents_md_competent_team_rule(agents_md):
    for role, text in agents_md.items():
        assert "flavor, not friction" in text, role


def test_engineer_git_flow(agents_md):
    text = agents_md["engineer"]
    for needle in ("{DEMO_REPO}", "fix/", "npm test", "git diff seeded-bug..HEAD", "infinity-fix.diff",
                   "gh pr create", "ARTIFACT: pr", "ARTIFACT: diff"):
        assert needle in text, needle


def test_sales_and_cfo_use_their_skills(agents_md):
    assert "{SKILLS_DIR}/pitch-deck/make_deck.py" in agents_md["sales_lead"]
    assert "{SKILLS_DIR}/excel-model/make_model.py" in agents_md["cfo"]


def test_skill_md_has_no_placeholders():
    # Agents read SKILL.md straight from disk, where nothing fills placeholders in.
    for skill_md in SKILLS.glob("*/SKILL.md"):
        assert not re.findall(r"\{[A-Z_]+\}", skill_md.read_text()), skill_md


@pytest.mark.parametrize("skill,script,ext", [("pitch-deck", "make_deck.py", "pptx"), ("excel-model", "make_model.py", "xlsx")])
def test_skill_md_example_spec_builds(tmp_path, skill, script, ext):
    example = re.search(r"```json\n(.*?)```", (SKILLS / skill / "SKILL.md").read_text(), re.S).group(1)
    spec = tmp_path / "spec.json"
    spec.write_text(example)
    result = run(SKILLS / skill / script, spec, tmp_path / f"out.{ext}")
    assert result.returncode == 0, result.stderr
